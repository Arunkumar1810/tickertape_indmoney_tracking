"""Read-only Angel One SmartAPI MCP server (stdio).

Exposes portfolio, funds, positions, order/trade books and market data.
Deliberately has NO order placement / modify / cancel / GTT tools.

Credentials are read from the repo-root .env (or the process environment):
  ANGEL_ONE_API_KEY, ANGEL_ONE_CLIENT_CODE, ANGEL_ONE_PASSWORD (PIN), ANGEL_ONE_TOTP_SECRET
Tokens are kept in memory only and are never returned from any tool.
"""
from __future__ import annotations

import datetime
import os
import re
import sys
import threading
from pathlib import Path
from typing import Any

import openpyxl
import pyotp
from dotenv import load_dotenv
from mcp.server.fastmcp import FastMCP
from SmartApi import SmartConnect

# Load <repo-root>/.env (server lives at <root>/.agents/mcp/angelone-server/server.py)
_ROOT = Path(__file__).resolve().parents[3]
load_dotenv(_ROOT / ".env")

_DATA_DIR = Path(os.environ.get("ANGELONE_DATA_DIR", Path(__file__).resolve().parent / "data"))

mcp = FastMCP("angelone")

ANGEL_SCRIP_MAP = {
    "ZYDUS LIFESCIENCES LTD": "ZYDUSLIFE",
    "HDFC BANK LTD": "HDFCBANK",
    "NATCO PHARMA LTD.": "NATCOPHARM",
    "TATA CONSULTANCY SERV LT": "TCS",
    "WIPRO LTD": "WIPRO",
    "THE SOUTH INDIAN BANK LTD": "SOUTHBANK",
    "INFOSYS LIMITED": "INFY",
    "ITC LTD": "ITC",
    "HCL TECHNOLOGIES LTD": "HCLTECH",
}


def _normalize_date(val: Any) -> str | None:
    if val is None:
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)):
        try:
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(val))).strftime("%Y-%m-%d")
        except Exception:
            return None
    s = str(val).strip().split(" ")[0]
    if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
        return s
    m = re.match(r"^(\d{1,2})[-/](\d{1,2})[-/](\d{4})$", s)
    if m:
        return f"{m.group(3)}-{m.group(2).zfill(2)}-{m.group(1).zfill(2)}"
    if re.match(r"^\d{5}$", s):
        try:
            return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=float(s))).strftime("%Y-%m-%d")
        except Exception:
            return None
    return None


def _parse_amount(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    clean = re.sub(r"[^\d.-]", "", str(val))
    try:
        return float(clean)
    except Exception:
        return 0.0


def _parse_tradebooks(year: int | None = None) -> list[dict[str, Any]]:
    """Scan multi-year tradebook files (.xlsx), extract all executed trades, deduplicate, and sort chronologically."""
    trade_files: list[Path] = []
    search_dirs = [_DATA_DIR / "tradebook", _DATA_DIR]
    for d in search_dirs:
        if d.is_dir():
            for f in d.glob("*.xlsx"):
                if f.name.startswith("~"):
                    continue
                if d.name == "tradebook" or "trade" in f.name.lower():
                    if f not in trade_files:
                        trade_files.append(f)

    trade_files.sort(key=lambda f: f.name)
    trades_by_key: dict[str, dict[str, Any]] = {}

    for file_path in trade_files:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_name = next((s for s in wb.sheetnames if "trade" in s.lower()), wb.sheetnames[0])
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            header_idx = -1
            for idx, r in enumerate(rows):
                lower_strs = [str(c).strip().lower() for c in r if c is not None]
                if any("scrip" in s or "symbol" in s for s in lower_strs) and any("buy" in s or "type" in s for s in lower_strs):
                    header_idx = idx
                    break

            if header_idx < 0:
                continue

            header_row = [str(c).strip().lower() if c is not None else "" for c in rows[header_idx]]

            def get_col(candidates: list[str]) -> int | None:
                for cand in candidates:
                    for i, h in enumerate(header_row):
                        if cand in h:
                            return i
                return None

            scrip_col = get_col(["scrip", "symbol", "contract"])
            type_col = get_col(["buy/sell", "trade type", "type", "transaction type"])
            buy_price_col = get_col(["buy price", "price", "rate", "avg price"])
            sell_price_col = get_col(["sell price", "price", "rate", "avg price"])
            qty_col = get_col(["quantity", "qty"])
            trade_id_col = get_col(["trade id", "trade_id", "order id", "order_id"])
            date_col = get_col(["trade date", "date", "time"])
            isin_col = get_col(["isin"])

            if scrip_col is None or type_col is None or date_col is None or qty_col is None:
                continue

            for r in rows[header_idx + 1:]:
                if not r or len(r) <= scrip_col or not r[scrip_col]:
                    continue
                raw_scrip = str(r[scrip_col]).strip()
                if raw_scrip.startswith("NOTE:"):
                    continue

                sym = ANGEL_SCRIP_MAP.get(raw_scrip, raw_scrip).upper()
                sym = re.sub(r"-(EQ|BE)$", "", sym).strip()
                if not sym:
                    continue

                raw_type = str(r[type_col] or "").strip().upper()
                if raw_type not in ("BUY", "SELL", "B", "S"):
                    continue
                is_buy = raw_type in ("BUY", "B")
                trade_type = "BUY" if is_buy else "SELL"

                p_col = buy_price_col if (is_buy and buy_price_col is not None) else sell_price_col
                price = _parse_amount(r[p_col]) if p_col is not None and len(r) > p_col else 0.0
                qty = _parse_amount(r[qty_col]) if len(r) > qty_col else 0.0
                if qty <= 0 or price <= 0:
                    continue

                date_val = _normalize_date(r[date_col]) if len(r) > date_col else None
                if not date_val:
                    continue
                if year and not date_val.startswith(str(year)):
                    continue

                trade_id = str(r[trade_id_col]).strip() if trade_id_col is not None and len(r) > trade_id_col and r[trade_id_col] else ""
                if not trade_id:
                    trade_id = f"{date_val}_{sym}_{trade_type}_{qty}_{price}"

                isin = str(r[isin_col]).strip().upper() if isin_col is not None and len(r) > isin_col and r[isin_col] else ""

                dedup_key = f"{trade_id}_{date_val}_{sym}_{trade_type}_{qty}_{price}"
                trades_by_key[dedup_key] = {
                    "tradeId": trade_id,
                    "symbol": sym,
                    "isin": isin,
                    "date": date_val,
                    "type": trade_type,
                    "quantity": qty,
                    "price": price,
                    "amount": round(qty * price, 2)
                }
        except Exception as e:
            _log(f"Warning: Failed to parse tradebook {file_path.name}: {e}")

    result = list(trades_by_key.values())
    result.sort(key=lambda t: (t["date"], t["tradeId"]))
    return result


def _parse_ledgers(year: int | None = None) -> list[dict[str, Any]]:
    """Scan multi-year account ledger files (.xlsx), extract bank fund pay-ins/payouts, deduplicate, and sort chronologically."""
    ledger_files: list[Path] = []
    search_dirs = [_DATA_DIR / "ledger", _DATA_DIR]
    for d in search_dirs:
        if d.is_dir():
            for f in d.glob("*.xlsx"):
                if f.name.startswith("~"):
                    continue
                if d.name == "ledger" or "ledger" in f.name.lower():
                    if f not in ledger_files:
                        ledger_files.append(f)

    ledger_files.sort(key=lambda f: f.name)
    cashflows_by_key: dict[str, dict[str, Any]] = {}

    for file_path in ledger_files:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_name = next((s for s in wb.sheetnames if "ledger" in s.lower()), wb.sheetnames[0])
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            header_idx = -1
            for idx, r in enumerate(rows):
                lower_strs = [str(c).strip().lower() for c in r if c is not None]
                if any("date" in s for s in lower_strs) and any("debit" in s for s in lower_strs) and any("credit" in s for s in lower_strs):
                    header_idx = idx
                    break

            if header_idx < 0:
                continue

            header_row = [str(c).strip().lower() if c is not None else "" for c in rows[header_idx]]

            def get_col(candidates: list[str]) -> int | None:
                for cand in candidates:
                    for i, h in enumerate(header_row):
                        if cand in h:
                            return i
                return None

            desc_col = get_col(["transaction", "particular", "description", "narrative"])
            date_col = get_col(["date", "posting date", "txn date"])
            voucher_col = get_col(["voucher"])
            debit_col = get_col(["debit"])
            credit_col = get_col(["credit"])

            if date_col is None:
                continue

            for r in rows[header_idx + 1:]:
                if not r or len(r) <= date_col or not r[date_col]:
                    continue
                date_val = _normalize_date(r[date_col])
                if not date_val:
                    continue
                if year and not date_val.startswith(str(year)):
                    continue

                desc = str(r[desc_col]).strip() if desc_col is not None and len(r) > desc_col and r[desc_col] else ""
                voucher = str(r[voucher_col]).strip() if voucher_col is not None and len(r) > voucher_col and r[voucher_col] else ""
                debit = _parse_amount(r[debit_col]) if debit_col is not None and len(r) > debit_col else 0.0
                credit = _parse_amount(r[credit_col]) if credit_col is not None and len(r) > credit_col else 0.0

                desc_lower = desc.lower()
                is_payin = (
                    "payin" in desc_lower or
                    "fund transfer" in desc_lower or
                    "funds added" in desc_lower or
                    "e-collect" in desc_lower or
                    "upi" in desc_lower or
                    "payment gateway" in desc_lower or
                    (credit > 0 and "dividend" not in desc_lower and "sell" not in desc_lower and "interest" not in desc_lower)
                )
                is_payout = (
                    "payout" in desc_lower or
                    "withdrawal" in desc_lower or
                    "funds withdrawn" in desc_lower or
                    ("settlement" in desc_lower and debit > 0)
                )
                is_dividend = "dividend" in desc_lower

                if is_payin and credit > 0:
                    key = f"PAYIN_{voucher}_{date_val}_{credit}"
                    cashflows_by_key[key] = {
                        "date": date_val,
                        "amount": -abs(credit),
                        "category": "PAYIN",
                        "description": desc or "Funds Added (Pay-in)",
                        "voucher": voucher
                    }
                elif is_payout and debit > 0:
                    key = f"PAYOUT_{voucher}_{date_val}_{debit}"
                    cashflows_by_key[key] = {
                        "date": date_val,
                        "amount": abs(debit),
                        "category": "PAYOUT",
                        "description": desc or "Funds Payout / Settlement",
                        "voucher": voucher
                    }
                elif is_dividend and credit > 0:
                    key = f"DIVIDEND_{voucher}_{date_val}_{credit}"
                    cashflows_by_key[key] = {
                        "date": date_val,
                        "amount": abs(credit),
                        "category": "DIVIDEND",
                        "description": desc or "Dividend Credit",
                        "voucher": voucher
                    }
        except Exception as e:
            _log(f"Warning: Failed to parse ledger {file_path.name}: {e}")

    result = list(cashflows_by_key.values())
    result.sort(key=lambda c: (c["date"], c.get("voucher", "")))
    return result


def _log(msg: str) -> None:
    # stdout is the MCP channel; all diagnostics go to stderr
    print(msg, file=sys.stderr, flush=True)


class _Session:
    """Lazily-created, auto-refreshing SmartAPI session."""

    def __init__(self) -> None:
        self._api: SmartConnect | None = None
        self._lock = threading.Lock()

    @staticmethod
    def _env(name: str) -> str:
        val = os.environ.get(name, "").strip()
        if not val:
            raise RuntimeError(f"Missing environment variable {name} (set it in the repo-root .env)")
        return val

    def _login(self) -> SmartConnect:
        api = SmartConnect(api_key=self._env("ANGEL_ONE_API_KEY"))
        totp = pyotp.TOTP(self._env("ANGEL_ONE_TOTP_SECRET")).now()
        res = api.generateSession(self._env("ANGEL_ONE_CLIENT_CODE"), self._env("ANGEL_ONE_PASSWORD"), totp)
        if not res or not res.get("status"):
            raise RuntimeError(f"Angel One login failed: {res.get('message') if res else 'no response'}")
        _log("Angel One session established")
        return api

    def get(self, force: bool = False) -> SmartConnect:
        with self._lock:
            if force or self._api is None:
                self._api = self._login()
            return self._api


_session = _Session()


def _call(fn_name: str, *args: Any, **kwargs: Any) -> Any:
    """Call a SmartConnect method, re-logging in once on an auth failure; return `data`."""
    last: Any = None
    for attempt in (0, 1):
        api = _session.get(force=attempt == 1)
        res = getattr(api, fn_name)(*args, **kwargs)
        last = res
        if isinstance(res, dict) and res.get("status"):
            return res.get("data")
        msg = str(res.get("message", "") if isinstance(res, dict) else res).lower()
        code = str(res.get("errorcode", "") if isinstance(res, dict) else "")
        if attempt == 0 and ("token" in msg or "session" in msg or code in {"AG8001", "AG8002", "AB8050"}):
            continue
        break
    raise RuntimeError(f"SmartAPI {fn_name} failed: {last.get('message') if isinstance(last, dict) else last}")


@mcp.tool()
def get_holdings() -> Any:
    """Equity holdings in the demat account (tradingsymbol, isin, quantity, averageprice, ltp, pnl)."""
    data = _call("holding")
    return data or []


@mcp.tool()
def get_all_holdings() -> Any:
    """All holdings with portfolio totals (holdings list + totalholding summary)."""
    return _call("allholding") or {}


@mcp.tool()
def get_positions() -> Any:
    """Net/day positions for today (intraday and carry-forward)."""
    return _call("position") or []


@mcp.tool()
def get_rms_limits() -> Any:
    """Funds and margin (RMS): net/available cash, collateral, utilised amounts."""
    return _call("rmsLimit") or {}


@mcp.tool()
def get_order_book() -> Any:
    """All orders placed today (read-only)."""
    return _call("orderBook") or []


@mcp.tool()
def get_trade_book(historical: bool = False, year: int | None = None) -> Any:
    """Trades executed. If historical=True, returns multi-year historical trades from tradebook statements; otherwise calls SmartAPI for today's trades."""
    if historical:
        return _parse_tradebooks(year=year)
    return _call("tradeBook") or []


@mcp.tool()
def get_historical_trades(year: int | None = None) -> list[dict[str, Any]]:
    """Multi-year historical trade executions parsed from annual Excel tradebook statements (.xlsx) in the MCP data directory, sorted chronologically. Supports optional year filter."""
    return _parse_tradebooks(year=year)


@mcp.tool()
def get_account_ledger(year: int | None = None) -> list[dict[str, Any]]:
    """Multi-year bank pay-ins, payouts, and settlements parsed from annual Angel One ledger statements in the MCP data directory, sorted chronologically. Supports optional year filter."""
    return _parse_ledgers(year=year)


@mcp.tool()
def get_ltp(exchange: str, tradingsymbol: str, symboltoken: str) -> Any:
    """Last traded price. Example: exchange='NSE', tradingsymbol='SBIN-EQ', symboltoken='3045'."""
    return _call("ltpData", exchange, tradingsymbol, symboltoken)


@mcp.tool()
def get_candle_data(exchange: str, symboltoken: str, interval: str, fromdate: str, todate: str) -> Any:
    """Historical OHLCV candles. interval e.g. ONE_DAY, ONE_HOUR; dates as 'YYYY-MM-DD HH:MM'."""
    return _call(
        "getCandleData",
        {
            "exchange": exchange,
            "symboltoken": symboltoken,
            "interval": interval,
            "fromdate": fromdate,
            "todate": todate,
        },
    )


if __name__ == "__main__":
    mcp.run(transport="stdio")
